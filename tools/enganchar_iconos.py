#!/usr/bin/env python3
"""Conecta sprites nuevos a sus .tres y actualiza el Excel.

Para cada godot/data/items/<id>.tres que NO tenga icono y cuyo
godot/art/sprites/items/<id>.png SÍ exista:

  - Añade ext_resource de la textura y línea icon = ExtResource(...)
  - Sube load_steps en 1
  - En Catalogo_Maestro.xlsx, hoja Items: cambia "⏳ FALTA SPRITE" → "✅ Obtainable"
    y columna Sprite a "static"

Uso:
    python3 enganchar_iconos.py           # engancha todo
    python3 enganchar_iconos.py --seco    # muestra qué haría sin escribir
    python3 enganchar_iconos.py id1 id2   # solo esos ids
"""

import os
import re
import shutil
import sys
from pathlib import Path

from openpyxl import load_workbook

# ── paths ────────────────────────────────────────────────────────────────
PROJECT = Path(__file__).resolve().parent.parent
DATA_DIR = PROJECT / "godot" / "data" / "items"
SPRITES_DIR = PROJECT / "godot" / "art" / "sprites" / "items"
EXCEL_PATH = PROJECT / "Catalogo_Maestro.xlsx"


def encontrar_pendientes(ids_filtro: list[str] | None = None) -> list[dict]:
    """Devuelve lista de dicts con id, tres_path, png_path para items pendientes."""
    pendientes = []

    for tres_path in sorted(DATA_DIR.glob("*.tres")):
        item_id = tres_path.stem  # sin .tres

        if ids_filtro and item_id not in ids_filtro:
            continue

        # Leer el .tres
        content = tres_path.read_text(encoding="utf-8")

        # ¿Ya tiene icono?
        if "icon = ExtResource" in content:
            continue

        # ¿Existe el PNG?
        png_path = SPRITES_DIR / f"{item_id}.png"
        if not png_path.exists():
            continue

        pendientes.append({
            "id": item_id,
            "tres_path": tres_path,
            "png_path": png_path,
            "content": content,
        })

    return pendientes


def enganchar_tres(item: dict) -> str:
    """Modifica el contenido del .tres para añadir el icono.

    Devuelve el nuevo contenido.
    """
    content = item["content"]
    item_id = item["id"]

    # 1. Subir load_steps en 1
    def bump_load_steps(match):
        n = int(match.group(1))
        return f'load_steps={n + 1}'

    content = re.sub(r'load_steps=(\d+)', bump_load_steps, content, count=1)

    # 2. Añadir ext_resource de la textura después del ext_resource del Script
    texture_line = (
        f'[ext_resource type="Texture2D" '
        f'path="res://art/sprites/items/{item_id}.png" id="2_icon"]'
    )

    # Insertar después de la línea del ext_resource del Script
    script_pattern = r'(\[ext_resource type="Script" path="res://scripts/data/item_data\.gd" id="1_id"\])'
    replacement = r'\1\n' + texture_line
    content = re.sub(script_pattern, replacement, content, count=1)

    # 3. Añadir icon = ExtResource("2_icon") después de script = ExtResource("1_id")
    script_assign = r'(script = ExtResource\("1_id"\))'
    icon_line = 'icon = ExtResource("2_icon")'
    replacement2 = r'\1\n' + icon_line
    content = re.sub(script_assign, replacement2, content, count=1)

    return content


def actualizar_excel(pendientes: list[dict], seco: bool = False) -> list[str]:
    """Actualiza el Excel: cambia estado y sprite de los items enganchados.

    Devuelve lista de mensajes de lo que se hizo/haría.
    """
    msgs = []

    if not pendientes:
        msgs.append("  Nada que actualizar en el Excel.")
        return msgs

    ids_a_actualizar = {p["id"] for p in pendientes}

    if seco:
        msgs.append(f"  [SECO] Haría .bak de {EXCEL_PATH.name}")
        msgs.append(
            f"  [SECO] Actualizaría {len(ids_a_actualizar)} filas en hoja Items"
        )
        return msgs

    # Backup
    bak_path = EXCEL_PATH.with_suffix(".xlsx.bak")
    shutil.copy2(EXCEL_PATH, bak_path)
    msgs.append(f"  .bak guardado en {bak_path.name}")

    wb = load_workbook(EXCEL_PATH)
    ws = wb["Items"]

    actualizados = 0
    for row in range(2, ws.max_row + 1):
        item_id = ws.cell(row=row, column=1).value  # Columna A: ID
        if item_id and str(item_id).strip() in ids_a_actualizar:
            estado = ws.cell(row=row, column=9).value  # Columna I: Estado
            if estado and "FALTA SPRITE" in str(estado):
                ws.cell(row=row, column=9).value = "✅ Obtainable"
                ws.cell(row=row, column=11).value = "static"  # Columna K: Sprite
                actualizados += 1

    wb.save(EXCEL_PATH)
    msgs.append(f"  {actualizados} filas actualizadas en hoja Items")
    return msgs


def main():
    seco = "--seco" in sys.argv
    ids_filtro = [a for a in sys.argv[1:] if not a.startswith("--")]

    if ids_filtro:
        print(f"Filtrando por ids: {ids_filtro}")

    pendientes = encontrar_pendientes(ids_filtro if ids_filtro else None)

    if not pendientes:
        print("No hay items pendientes de enganchar.")
        return

    print(f"\nItems pendientes de enganchar: {len(pendientes)}")
    for p in pendientes:
        print(f"  ▸ {p['id']}")

    if seco:
        print("\n⚠️  MODO SECO — no se escribe nada.\n")
        for p in pendientes:
            nuevo = enganchar_tres(p)
            print(f"  [{p['id']}.tres] load_steps subiría +1, añadiría ext_resource + icon")
        msgs = actualizar_excel(pendientes, seco=True)
        for m in msgs:
            print(m)
        print(f"\n  {len(pendientes)} items se engancharían.")
        return

    # ── ESCRIBIR ──
    print("\nEnganchando...")
    for p in pendientes:
        nuevo = enganchar_tres(p)
        p["tres_path"].write_text(nuevo, encoding="utf-8")
        print(f"  ✓ {p['id']}.tres")

    msgs = actualizar_excel(pendientes, seco=False)
    for m in msgs:
        print(m)

    print(f"\n✓ {len(pendientes)} items enganchados.")


if __name__ == "__main__":
    main()
